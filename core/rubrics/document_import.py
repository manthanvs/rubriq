"""Read a rubric out of a document the department already has.

A guide may keep her rubric in a spreadsheet, a Word file, or a PDF of the
syllabus. All three are accepted here.

**Which route a file takes is decided by what the file is, not by preference.**

A spreadsheet laid out as a rubric table is parsed exactly, by
:mod:`core.rubrics.excel_import` — no model involved, because the shape is
regular and a parser is faster, free, offline, and cannot reword a criterion.

Everything else goes to the model, and this is where it earns its place: a
syllabus PDF or a Word table has no fixed shape, and reading prose into
structure is the thing a parser cannot do and a language model can.

What the model is allowed to do is deliberately narrow. It extracts; it does
not judge, score, or invent. Its output is validated against a schema before
anything is written, and every criterion it returns must carry a title — so a
hallucinated rubric fails loudly rather than arriving as plausible nonsense.
"""

from __future__ import annotations

import json
import re
from decimal import Decimal, InvalidOperation

from core.ai.provider import AIUnavailable, LLMProvider
from core.errors import ValidationError
from core.rubrics.excel_import import (
    ParsedCriterion,
    ParsedLevel,
    ParsedReview,
    parse_rubric_workbook,
)
from core.submissions.extract import extract_text

#: Everything a rubric may arrive as.
ALLOWED_EXTENSIONS = (".xlsx", ".pdf", ".docx", ".txt", ".md")

#: Enough of a document to hold a rubric, without paying for a whole syllabus.
#: A rubric that has not appeared in sixty thousand characters is not going to.
MAX_CHARS = 60_000

PROMPT_VERSION = "rubric-extract-v1"

SYSTEM_PROMPT = """You extract assessment rubrics from documents.

Return ONLY JSON, no prose and no code fences, shaped exactly like this:

{
  "reviews": [
    {
      "index": 1,
      "title": "short title of the review or assessment",
      "max_marks": 25,
      "criteria": [
        {
          "title": "name of the criterion",
          "max_score": 5,
          "levels": [
            {"label": "Excellent", "score": 5, "descriptor": "what earns it"},
            {"label": "Poor", "score": 0, "descriptor": "what earns it"}
          ]
        }
      ]
    }
  ]
}

Rules you must follow:
- Copy titles and descriptors from the document. Do not reword or improve them.
- If a level is given as a range such as "Satisfactory (2-3)", its score is the
  top of that range.
- max_marks is the total for that review. If the document states it, use it. If
  it does not, use the sum of the criteria's max_score values.
- Include every criterion you find, in the order the document gives them.
- If the document contains no assessment rubric at all, return {"reviews": []}.
  Never invent criteria to fill the response."""


def _user_prompt(text: str) -> str:
    return (
        "Extract every assessment rubric from the document below.\n\n"
        "--- DOCUMENT ---\n" + text[:MAX_CHARS] + "\n--- END ---"
    )


def _decimal(value: object, *, field: str) -> Decimal:
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise ValidationError(
            f"The model returned {value!r} for {field}, which is not a number."
        ) from exc


def _strip_fences(raw: str) -> str:
    """Models add code fences even when told not to."""
    text = raw.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    return text.strip()


def parse_ai_response(raw: str) -> tuple[ParsedReview, ...]:
    """Validate a model response into reviews, or say why it cannot be."""
    try:
        payload = json.loads(_strip_fences(raw))
    except json.JSONDecodeError as exc:
        raise ValidationError(
            "The model did not return readable JSON. Try the file again, or "
            "save the rubric as a spreadsheet, which is read without a model."
        ) from exc

    if not isinstance(payload, dict) or not isinstance(payload.get("reviews"), list):
        raise ValidationError("The model's response did not contain a review list.")

    reviews: list[ParsedReview] = []
    for position, item in enumerate(payload["reviews"], 1):
        if not isinstance(item, dict):
            continue

        raw_criteria = item.get("criteria")
        if not isinstance(raw_criteria, list) or not raw_criteria:
            continue

        criteria_titles: list[tuple[str, tuple[tuple[str, str], ...], Decimal]] = []
        labels: dict[str, Decimal] = {}

        for entry in raw_criteria:
            if not isinstance(entry, dict):
                continue
            title = str(entry.get("title") or "").strip()
            if not title:
                # A criterion with no name cannot be marked against and is the
                # clearest sign the model was guessing.
                raise ValidationError(
                    "The model returned a criterion with no title. Nothing was "
                    "imported — check the document actually contains a rubric."
                )

            max_score = _decimal(entry.get("max_score", 5), field="max_score")
            descriptors: list[tuple[str, str]] = []
            for level in entry.get("levels") or []:
                if not isinstance(level, dict):
                    continue
                label = str(level.get("label") or "").strip()
                if not label:
                    continue
                descriptors.append((label, str(level.get("descriptor") or "").strip()))
                labels.setdefault(label, _decimal(level.get("score", 0), field="score"))

            criteria_titles.append((title, tuple(descriptors), max_score))

        if not criteria_titles:
            continue

        share = (Decimal("100") / len(criteria_titles)).quantize(Decimal("0.01"))
        weights = [share] * len(criteria_titles)
        weights[-1] = Decimal("100") - share * (len(criteria_titles) - 1)

        declared = item.get("max_marks")
        max_marks = (
            _decimal(declared, field="max_marks")
            if declared is not None
            else sum((c[2] for c in criteria_titles), Decimal("0"))
        )

        reviews.append(
            ParsedReview(
                index=int(item.get("index") or position),
                title=str(item.get("title") or f"Review {position}").strip(),
                max_marks=max_marks,
                levels=tuple(
                    ParsedLevel(label=label, score=score)
                    for label, score in labels.items()
                ),
                criteria=tuple(
                    ParsedCriterion(
                        code=f"C{number}",
                        title=title,
                        descriptors=descriptors,
                        max_score=max_score,
                        weight=weight,
                    )
                    for number, (title, descriptors, max_score), weight in zip(
                        range(1, len(criteria_titles) + 1),
                        criteria_titles,
                        weights,
                        strict=True,
                    )
                ),
            )
        )

    if not reviews:
        raise ValidationError(
            "No rubric was found in that document. It may not contain one, or "
            "the assessment table may not be readable as text — a scanned PDF "
            "has no text to read."
        )
    return tuple(reviews)


def read_rubric_document(
    filename: str,
    data: bytes,
    *,
    provider: LLMProvider | None = None,
) -> tuple[tuple[ParsedReview, ...], str]:
    """Read a rubric from any supported document.

    Returns the reviews and how they were obtained, so a page can tell the
    guide whether a model was involved in what she is about to publish.
    """
    extension = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if extension not in ALLOWED_EXTENSIONS:
        raise ValidationError(
            f"{extension or 'That file'} is not supported. Upload a .xlsx, "
            ".pdf, .docx or .txt file."
        )

    if extension == ".xlsx":
        try:
            return parse_rubric_workbook(data), "read directly from the spreadsheet"
        except ValidationError:
            # A workbook that is not laid out as a rubric table still has text
            # in it, so fall through and let the model read it.
            text = _workbook_text(data)
    else:
        text = extract_text(filename, data).text

    if not text.strip():
        raise ValidationError(
            "No text could be read from that file. A scanned PDF is an image, "
            "and this does not read images."
        )

    if provider is None:
        raise AIUnavailable(
            "Reading a rubric out of a PDF or Word file needs an AI provider "
            "configured. A spreadsheet laid out as a rubric table is read "
            "without one."
        )

    raw = provider.complete(system=SYSTEM_PROMPT, user=_user_prompt(text))
    return parse_ai_response(raw), f"read by {provider.name}"


def _workbook_text(data: bytes) -> str:
    """Every cell of a workbook as plain text, for the model to read."""
    import io

    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ValidationError("That file could not be opened as a workbook.") from exc

    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
    workbook.close()
    return "\n".join(lines)
