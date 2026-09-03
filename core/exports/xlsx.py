"""XLSX export — a real spreadsheet, not a CSV with the wrong extension (§3).

Writes exactly ``bundle.cells()`` on the first sheet, so the numbers are the
same objects the TSV writes (fix item 11). Styling is applied *around* those
values; nothing here reformats one.

A second sheet carries evidence and rationale per criterion, as §7 requires —
that is what makes an exported mark defensible when a student asks where it
came from.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.exports.rows import ESTIMATE_MARKER, ExportBundle

HEADER_FILL = PatternFill("solid", fgColor="2E5AAC")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ESTIMATE_FILL = PatternFill("solid", fgColor="FFF3CD")
ABSENT_FONT = Font(bold=True, color="B00020")

MAX_COLUMN_WIDTH = 42


def _autosize(worksheet, grid: list[list[str]]) -> None:
    for index in range(len(grid[0])):
        widest = max((len(str(row[index])) for row in grid), default=10)
        letter = get_column_letter(index + 1)
        worksheet.column_dimensions[letter].width = min(widest + 2, MAX_COLUMN_WIDTH)


def to_xlsx(bundle: ExportBundle) -> bytes:
    """Build the workbook and return its bytes."""
    workbook = Workbook()

    scores = workbook.active
    scores.title = "Scores"

    grid = bundle.cells()
    for row in grid:
        scores.append(row)

    for cell in scores[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    # Frozen header row, so a cohort of thirty stays readable when scrolled.
    scores.freeze_panes = "A2"

    status_column = bundle.headers.index("Status") + 1
    final_column = bundle.headers.index("Final") + 1

    for row_index in range(2, len(grid) + 1):
        status_value = str(scores.cell(row=row_index, column=status_column).value or "")
        final_value = str(scores.cell(row=row_index, column=final_column).value or "")

        # An unapproved row is tinted as well as labelled: invariant #1 should
        # be visible at a glance, not only on close reading.
        if ESTIMATE_MARKER in status_value:
            for column in range(1, len(bundle.headers) + 1):
                scores.cell(row=row_index, column=column).fill = ESTIMATE_FILL

        if final_value == "ABSENT":
            scores.cell(row=row_index, column=final_column).font = ABSENT_FONT

    _autosize(scores, grid)

    # --- sheet two: evidence -------------------------------------------
    evidence = workbook.create_sheet("Evidence")
    evidence_headers = [
        "PRN",
        "Name",
        "Criterion",
        "Verdict",
        "Score",
        "Evidence",
        "Rationale",
    ]
    evidence.append(evidence_headers)

    for entry in bundle.evidence:
        evidence.append([entry.get(header, "") for header in evidence_headers])

    for cell in evidence[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    evidence.freeze_panes = "A2"

    for column, width in zip(
        "ABCDEFG", (14, 24, 10, 14, 8, MAX_COLUMN_WIDTH, MAX_COLUMN_WIDTH), strict=False
    ):
        evidence.column_dimensions[column].width = width

    for row in evidence.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # --- sheet three: what this file is --------------------------------
    about = workbook.create_sheet("About")
    for label, value in (
        ("Subject", f"{bundle.subject_code} — {bundle.subject_name}"),
        ("Milestone", f"Review {bundle.milestone_index} — {bundle.milestone_title}"),
        ("Max marks", bundle.max_marks),
        ("Generated (IST)", bundle.filename_stem.rsplit("_", 1)[-1]),
        ("Rows", str(len(bundle.rows))),
        (
            "Note",
            "Rows marked "
            + ESTIMATE_MARKER
            + " have not been approved by a faculty member and are not final.",
        ),
    ):
        about.append([label, value])

    about.column_dimensions["A"].width = 18
    about.column_dimensions["B"].width = 70
    for row in about.iter_rows():
        row[0].font = Font(bold=True)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
