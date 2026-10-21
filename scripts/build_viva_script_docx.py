"""Build a printable version of a presenting document.

Defaults to the viva script; ``--source`` points it at another, such as
``docs/program-flow.md``. Both are documents you hold while talking, and
both are generated rather than kept as a second hand-edited copy.

Generated *from* ``docs/viva-script.md`` rather than written again here. The
script is a thing you hold while talking, so it has to print; but a second
hand-maintained copy is how two versions of the same instructions end up
disagreeing on the morning you need them.

The Markdown subset is only what that file uses — headings, tables, block
quotes, bullets, rules and inline bold/code/links. It is not a general
converter and does not try to be.

    make viva
    python scripts/build_viva_script_docx.py
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

from scripts.build_synopsis_docx import (  # noqa: E402
    BODY_PT,
    FONT,
    LINE_SPACING,
    centred,
    heading,
    page_break,
    page_number_footer,
    set_styles,
    two_column_table,
)

SOURCE = REPO_ROOT / "docs" / "viva-script.md"
DEFAULT_OUT = REPO_ROOT / "docs" / "report" / "RubriQ_Viva_Script.docx"

#: Column widths per table, by how many columns it has. The script's tables
#: are narrow-then-wide in every case.
WIDTHS = {2: (2.1, 4.2), 3: (0.4, 2.9, 3.0)}

_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _runs(paragraph, text: str, *, italic: bool = False, size: int = BODY_PT) -> None:
    """Write text, honouring **bold**, ``code`` and [links](...)."""
    for piece in _INLINE.split(text):
        if not piece:
            continue
        bold = code = emph = False
        if piece.startswith("**") and piece.endswith("**"):
            piece, bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*"):
            piece, emph = piece[1:-1], True
        elif piece.startswith("`") and piece.endswith("`"):
            piece, code = piece[1:-1], True
        elif piece.startswith("["):
            label, _, url = piece[1:].partition("](")
            piece = f"{label} ({url.rstrip(')')})" if "http" in url else label

        run = paragraph.add_run(piece)
        run.font.name = "Consolas" if code else FONT
        run.font.size = Pt(size - 1 if code else size)
        run.bold = bold
        run.italic = italic or emph


def _para(document, text: str, *, italic=False, indent=0.0, space_after=6) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = LINE_SPACING
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.left_indent = Inches(indent)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _runs(paragraph, text, italic=italic)


def _quote(document, lines: list[str]) -> None:
    """A spoken line. Indented and italic — it is meant to be said, not read."""
    text = " ".join(lines).strip()
    if text:
        _para(document, text, italic=True, indent=0.35, space_after=8)


def _bullet(document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.line_spacing = LINE_SPACING
    paragraph.paragraph_format.space_after = Pt(2)
    _runs(paragraph, text)


def _plain(text: str) -> str:
    """Strip emphasis markers. ``heading()`` takes text, not Markdown."""
    return text.replace("**", "").replace("*", "").strip()


def _code_block(document, lines: list[str]) -> None:
    """A fenced block, kept as drawn.

    Diagrams made of text are the reason this exists: joining their lines the
    way prose is joined turns an arrow diagram into a sentence.
    """
    for raw in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Inches(0.2)
        run = paragraph.add_run(raw)
        run.font.name = "Consolas"
        run.font.size = Pt(BODY_PT - 3)
    document.add_paragraph()


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build(out_path: Path = DEFAULT_OUT, source: Path = SOURCE) -> Path:
    lines = source.read_text(encoding="utf-8").splitlines()

    document = Document()
    set_styles(document)
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    page_number_footer(section)

    quote: list[str] = []
    table: list[list[str]] = []
    body: list[str] = []
    body_is_item = False
    code: list[str] | None = None
    first_section = True

    def flush() -> None:
        """Emit whatever has been accumulating.

        Body text is buffered because the source is hard-wrapped at about
        seventy-six columns. Treating each wrapped line as its own paragraph
        is what made the first printed draft ragged, and it also split a bold
        span across two lines so its markers survived into the document.
        """
        nonlocal quote, table, body, body_is_item
        if body:
            joined = " ".join(body)
            _bullet(document, joined) if body_is_item else _para(document, joined)
            body, body_is_item = [], False
        if quote:
            _quote(document, quote)
            quote = []
        if table:
            header, *rows = table
            widths = WIDTHS.get(len(header), (6.3,) * len(header))
            two_column_table(document, rows, headers=header, widths=widths)
            table = []

    for raw in lines:
        line = raw.rstrip()

        if line.startswith("```"):
            if code is None:
                flush()
                code = []
            else:
                _code_block(document, code)
                code = None
            continue
        if code is not None:
            code.append(raw)
            continue

        if line.startswith("> "):
            if body:
                flush()
            quote.append(line[2:])
            continue
        if line.startswith("|") and set(line) <= set("|-: "):
            continue  # the ---|--- separator row
        if line.startswith("|"):
            if not table:  # a new table starts: close whatever preceded it
                flush()
            table.append(_split_row(line))
            continue

        # Anything that is not one of these is a continuation of the
        # paragraph being built, so it must not flush — flushing here is what
        # turned every hard-wrapped line into its own paragraph.
        structural = (
            not line
            or line.startswith(("---", "# ", "## ", "### ", "- ", "* "))
            or re.match(r"^\d+\. ", line) is not None
        )
        if not structural:
            # A wrapped continuation of whatever is open — a paragraph or a
            # list item. Emitting it separately is what split a bold span
            # across two paragraphs and left its markers showing.
            if quote or table:
                flush()
            body.append(line)
            continue

        flush()

        if not line or line.startswith("---"):
            continue
        if line.startswith("# "):
            centred(document, _plain(line[2:]), size=20, bold=True, space_after=4)
            centred(
                document,
                "What to click, and what to say. Spoken lines are in italics.",
                size=11,
                space_after=18,
            )
        elif line.startswith("## "):
            if not first_section:
                page_break(document)
            first_section = False
            heading(document, _plain(line[3:]), level=1)
        elif line.startswith("### "):
            heading(document, _plain(line[4:]), level=2)
        elif line.startswith(("- ", "* ")):
            body, body_is_item = [line[2:]], True
        elif re.match(r"^\d+\. ", line):
            body, body_is_item = [re.sub(r"^\d+\. ", "", line)], True
        else:  # pragma: no cover - every line is classified above
            body.append(line)

    flush()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(out_path)
    return out_path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--source", type=Path, default=SOURCE)
    args = parser.parse_args()
    out = args.out if args.out.is_absolute() else REPO_ROOT / args.out
    source = args.source if args.source.is_absolute() else REPO_ROOT / args.source
    print(f"Wrote {build(out, source).relative_to(REPO_ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
