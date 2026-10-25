"""Reading a rubric out of a PDF, a Word file, or an odd spreadsheet.

No live model. A stub returns the response, so every branch — good JSON, code
fences, a hallucinated criterion, an empty result — is exercised offline and
deterministically. The one thing a live call would add is proof that a
particular vendor is reachable today, which is a fact about a key rather than
about this code.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from core.ai.provider import AIUnavailable, StubProvider
from core.errors import ValidationError
from core.rubrics.document_import import (
    parse_ai_response,
    read_rubric_document,
)

FIXTURES = Path(__file__).resolve().parents[3] / "tests" / "fixtures"

GOOD = """{
  "reviews": [
    {
      "index": 1,
      "title": "Project Introduction",
      "max_marks": 25,
      "criteria": [
        {"title": "Problem Definition", "max_score": 5,
         "levels": [{"label": "Excellent", "score": 5, "descriptor": "clear"},
                    {"label": "Poor", "score": 0, "descriptor": "missing"}]},
        {"title": "System Design", "max_score": 5,
         "levels": [{"label": "Excellent", "score": 5, "descriptor": "detailed"}]}
      ]
    }
  ]
}"""


def _docx(lines: list[str]) -> bytes:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# -- the spreadsheet route: no model involved ----------------------------


def test_a_rubric_spreadsheet_is_parsed_not_modelled() -> None:
    """The regular case must never cost an API call."""
    data = (FIXTURES / "department_rubric.xlsx").read_bytes()
    provider = StubProvider("{}")

    reviews, how = read_rubric_document("rubric.xlsx", data, provider=provider)

    assert len(reviews) == 2
    assert how == "read directly from the spreadsheet"


def test_a_spreadsheet_that_is_not_a_rubric_table_falls_through_to_the_model() -> None:
    """Cells still hold text, so the model gets a chance before giving up."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Our rubric, written out as a paragraph"])
    sheet.append(["Problem Definition is worth 5 marks"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    reviews, how = read_rubric_document(
        "notes.xlsx", buffer.getvalue(), provider=StubProvider(GOOD)
    )

    assert [c.title for c in reviews[0].criteria] == [
        "Problem Definition",
        "System Design",
    ]
    assert "stub" in how.lower() or "read by" in how


# -- the model route -----------------------------------------------------


def test_a_word_file_is_read_by_the_model() -> None:
    data = _docx(["Review 1", "Problem Definition 5 marks", "System Design 5 marks"])

    reviews, how = read_rubric_document("rubric.docx", data, provider=StubProvider(GOOD))

    assert len(reviews) == 1
    assert reviews[0].max_marks == 25
    assert how.startswith("read by")


def test_without_a_provider_it_says_so_rather_than_failing_obscurely() -> None:
    data = _docx(["Review 1", "Problem Definition"])

    with pytest.raises(AIUnavailable, match="needs an AI provider"):
        read_rubric_document("rubric.docx", data, provider=None)


def test_an_unsupported_extension_is_refused() -> None:
    with pytest.raises(ValidationError, match="not supported"):
        read_rubric_document("rubric.rtf", b"anything", provider=StubProvider(GOOD))


def test_a_file_with_no_readable_text_is_refused() -> None:
    """A scanned PDF is an image; there is nothing to read."""
    with pytest.raises(ValidationError, match="No text could be read"):
        read_rubric_document("empty.txt", b"   ", provider=StubProvider(GOOD))


# -- what the model returns ----------------------------------------------


def test_code_fences_are_tolerated() -> None:
    """Models add fences even when the prompt forbids them."""
    reviews = parse_ai_response("```json\n" + GOOD + "\n```")
    assert len(reviews[0].criteria) == 2


def test_weights_sum_to_one_hundred() -> None:
    (review,) = parse_ai_response(GOOD)
    assert sum(c.weight for c in review.criteria) == 100


def test_max_marks_falls_back_to_the_sum_of_the_criteria() -> None:
    raw = """{"reviews": [{"index": 1, "title": "R", "criteria": [
        {"title": "One", "max_score": 5, "levels": []},
        {"title": "Two", "max_score": 5, "levels": []}]}]}"""
    (review,) = parse_ai_response(raw)
    assert review.max_marks == 10


def test_a_criterion_with_no_title_is_refused() -> None:
    """The clearest sign the model was guessing. Nothing is imported."""
    raw = """{"reviews": [{"index": 1, "title": "R", "criteria": [
        {"title": "", "max_score": 5, "levels": []}]}]}"""

    with pytest.raises(ValidationError, match="no title"):
        parse_ai_response(raw)


def test_an_empty_review_list_is_refused_with_a_reason() -> None:
    """The prompt tells it to return nothing rather than invent — so honour it."""
    with pytest.raises(ValidationError, match="No rubric was found"):
        parse_ai_response('{"reviews": []}')


def test_unreadable_json_points_at_the_spreadsheet_route() -> None:
    with pytest.raises(ValidationError, match="spreadsheet"):
        parse_ai_response("I'm afraid I can't do that.")
