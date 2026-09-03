"""Fix item 11 — export consistency.

The proof it asks for: *"a parity test asserting the TSV cell grid equals the
XLSX cell grid, value for value."* That is :class:`TestParity` below.

The other named failures each get a test too: an absent student exporting as
``0``, an unapproved estimate exporting as though it were final, and the
filename drifting from §7.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from core.clock import IST
from core.db.engine import session_scope
from core.exports.rows import ESTIMATE_MARKER, build_export_rows
from core.exports.tsv import to_tsv
from core.exports.xlsx import to_xlsx
from core.scoring.enums import Verdict
from core.scoring.sheets import approve_sheet, reinstate, save_manual_scores

GENERATED_AT = datetime(2026, 9, 3, 10, 30, tzinfo=IST)
FULL = {"C1": (10, Verdict.FOLLOWED), "C2": (10, Verdict.FOLLOWED)}


def _bundle(db_factory, world, milestone_id):
    with session_scope(db_factory) as session:
        return build_export_rows(
            world.faculty_a,
            session,
            milestone_id=milestone_id,
            generated_at=GENERATED_AT,
        )


def _xlsx_cells(data: bytes) -> list[list[str]]:
    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook["Scores"]
    return [
        ["" if cell is None else str(cell) for cell in row]
        for row in sheet.iter_rows(values_only=True)
    ]


def _tsv_cells(text: str) -> list[list[str]]:
    return [line.split("\t") for line in text.split("\n")]


@pytest.fixture
def scored(db_factory, world, graded):
    """One scored, unapproved submission."""
    with session_scope(db_factory) as session:
        return save_manual_scores(
            world.faculty_a,
            session,
            submission_id=graded.submission_id,
            scores=FULL,
        )


class TestParity:
    def test_the_two_exports_agree_cell_for_cell(
        self, db_factory, world, graded, scored
    ) -> None:
        bundle = _bundle(db_factory, world, graded.milestone_id)

        assert _tsv_cells(to_tsv(bundle)) == _xlsx_cells(to_xlsx(bundle))

    def test_they_still_agree_once_approved(
        self, db_factory, world, graded, scored
    ) -> None:
        with session_scope(db_factory) as session:
            approve_sheet(world.faculty_a, session, score_sheet_id=scored.id)

        bundle = _bundle(db_factory, world, graded.milestone_id)

        assert _tsv_cells(to_tsv(bundle)) == _xlsx_cells(to_xlsx(bundle))

    def test_they_agree_when_a_student_has_not_submitted(
        self, db_factory, world, graded, scored
    ) -> None:
        """The non-submitter row is the one most likely to be built twice."""
        bundle = _bundle(db_factory, world, graded.milestone_id)

        assert len(bundle.rows) >= 1
        assert _tsv_cells(to_tsv(bundle)) == _xlsx_cells(to_xlsx(bundle))

    def test_both_start_with_the_same_header(
        self, db_factory, world, graded, scored
    ) -> None:
        bundle = _bundle(db_factory, world, graded.milestone_id)

        assert _tsv_cells(to_tsv(bundle))[0] == list(bundle.headers)
        assert _xlsx_cells(to_xlsx(bundle))[0] == list(bundle.headers)


class TestApprovalIsVisible:
    def test_an_unapproved_row_is_marked_as_an_estimate(
        self, db_factory, world, graded, scored
    ) -> None:
        """Invariant #1 must survive to the last step of the pipeline."""
        bundle = _bundle(db_factory, world, graded.milestone_id)
        text = to_tsv(bundle)

        assert ESTIMATE_MARKER in text

    def test_an_approved_row_carries_no_estimate_marker(
        self, db_factory, world, graded, scored
    ) -> None:
        with session_scope(db_factory) as session:
            approve_sheet(world.faculty_a, session, score_sheet_id=scored.id)

        bundle = _bundle(db_factory, world, graded.milestone_id)
        row = next(r for r in bundle.rows if r.email == world.student_1.email)

        assert ESTIMATE_MARKER not in row.status
        assert row.approved_by == world.faculty_a.email
        assert row.approved_at

    def test_the_approver_reaches_the_spreadsheet(
        self, db_factory, world, graded, scored
    ) -> None:
        with session_scope(db_factory) as session:
            approve_sheet(world.faculty_a, session, score_sheet_id=scored.id)

        bundle = _bundle(db_factory, world, graded.milestone_id)

        assert "Approved By" in bundle.headers
        assert world.faculty_a.email in to_tsv(bundle)


class TestAbsentNeverExportsAsANumber:
    def test_absent_renders_as_the_word(self, db_factory, world, graded) -> None:
        """Fix item 10, at the export boundary: ABSENT is not 0."""
        from core.db.models import ReviewMilestone, Submission

        # Push the submission six days past the deadline.
        with session_scope(db_factory) as session:
            submission = session.get(Submission, graded.submission_id)
            milestone = session.get(ReviewMilestone, graded.milestone_id)
            submission.submitted_at = milestone.due_at + timedelta(days=6)

        with session_scope(db_factory) as session:
            sheet = save_manual_scores(
                world.faculty_a,
                session,
                submission_id=graded.submission_id,
                scores=FULL,
            )

        assert sheet.is_absent

        bundle = _bundle(db_factory, world, graded.milestone_id)
        row = next(r for r in bundle.rows if r.email == world.student_1.email)

        assert row.final_total == "ABSENT"
        assert "ABSENT" in to_tsv(bundle)

        cells = _xlsx_cells(to_xlsx(bundle))
        final_index = list(bundle.headers).index("Final")
        assert any(row[final_index] == "ABSENT" for row in cells[1:])

    def test_a_reinstated_student_scores_again(self, db_factory, world, graded) -> None:
        from core.db.models import ReviewMilestone, Submission

        with session_scope(db_factory) as session:
            submission = session.get(Submission, graded.submission_id)
            milestone = session.get(ReviewMilestone, graded.milestone_id)
            submission.submitted_at = milestone.due_at + timedelta(days=6)

        with session_scope(db_factory) as session:
            sheet = save_manual_scores(
                world.faculty_a,
                session,
                submission_id=graded.submission_id,
                scores=FULL,
            )

        with session_scope(db_factory) as session:
            after = reinstate(
                world.faculty_a,
                session,
                score_sheet_id=sheet.id,
                reason="Hospitalised — certificate on file.",
            )

        assert after.final_total == Decimal("25.00")
        assert after.is_absent is False


class TestFilename:
    def test_it_matches_section_7_exactly(
        self, db_factory, world, graded, scored
    ) -> None:
        bundle = _bundle(db_factory, world, graded.milestone_id)

        assert bundle.filename_stem == "RubriQ_MCA33EL03_Review1_20260903"

    def test_the_date_is_ist_not_utc(self, db_factory, world, graded, scored) -> None:
        """00:30 IST on the 3rd is 19:00 UTC on the 2nd — the file is dated IST."""
        with session_scope(db_factory) as session:
            bundle = build_export_rows(
                world.faculty_a,
                session,
                milestone_id=graded.milestone_id,
                generated_at=datetime(2026, 9, 3, 0, 30, tzinfo=IST),
            )

        assert bundle.filename_stem.endswith("20260903")


class TestWorkbookShape:
    def test_it_has_the_evidence_sheet_section_7_requires(
        self, db_factory, world, graded, scored
    ) -> None:
        bundle = _bundle(db_factory, world, graded.milestone_id)
        workbook = load_workbook(io.BytesIO(to_xlsx(bundle)))

        assert "Scores" in workbook.sheetnames
        assert "Evidence" in workbook.sheetnames

    def test_the_header_row_is_frozen(self, db_factory, world, graded, scored) -> None:
        bundle = _bundle(db_factory, world, graded.milestone_id)
        workbook = load_workbook(io.BytesIO(to_xlsx(bundle)))

        assert workbook["Scores"].freeze_panes == "A2"

    def test_a_tab_inside_a_note_cannot_break_the_tsv_columns(
        self, db_factory, world, graded, scored
    ) -> None:
        with session_scope(db_factory) as session:
            approve_sheet(
                world.faculty_a,
                session,
                score_sheet_id=scored.id,
                faculty_note="Line one\twith a tab\nand a newline",
            )

        bundle = _bundle(db_factory, world, graded.milestone_id)
        rows = _tsv_cells(to_tsv(bundle))
        widths = {len(row) for row in rows}

        assert len(widths) == 1, f"ragged TSV: rows have widths {widths}"
